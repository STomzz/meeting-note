"""纪要 / 导出接口测试。"""

from __future__ import annotations

from tests.conftest import DEFAULT_MINUTES


async def test_create_minutes(client, fake_provider):
    response = await client.post(
        "/v1/minutes",
        json={"transcript": "张三：下周三灰度。李四：我来写公告。", "title": "产品周会"},
        headers={"Authorization": "Bearer user-key"},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["minutes"]["title"] == "产品周会"
    assert payload["markdown"].startswith("# 产品周会")
    assert payload["model"]
    assert payload["usage"]["total_tokens"] == 42


async def test_create_minutes_docx(client):
    response = await client.post(
        "/v1/minutes/docx",
        json={"transcript": "张三：下周三灰度。", "title": "产品周会"},
        headers={"Authorization": "Bearer user-key"},
    )
    assert response.status_code == 200
    assert response.content[:2] == b"PK"  # docx 是 zip 容器
    disposition = response.headers["content-disposition"]
    assert "UTF-8''" in disposition


async def test_export_docx(client):
    response = await client.post(
        "/v1/export/docx",
        json={"minutes": DEFAULT_MINUTES, "include_transcript": True, "transcript": "全文内容"},
        headers={"Authorization": "Bearer user-key"},
    )
    assert response.status_code == 200
    assert response.content[:2] == b"PK"


async def test_minutes_validation_error(client):
    response = await client.post(
        "/v1/minutes",
        json={"title": "缺少 transcript"},
        headers={"Authorization": "Bearer user-key"},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "validation_error"


async def test_docx_contains_expected_text(client):
    """导出的 docx 能被解析且包含关键内容。"""

    import io

    from docx import Document

    response = await client.post(
        "/v1/export/docx",
        json={"minutes": DEFAULT_MINUTES},
        headers={"Authorization": "Bearer user-key"},
    )
    document = Document(io.BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "产品周会" in text
    assert "会议摘要" in text
    assert "准备发布公告" in table_text
