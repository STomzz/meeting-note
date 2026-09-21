"""Word（.docx）导出：结构化纪要约 .docx 字节流（纯函数，便于单测）。"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.schemas.minutes import Minutes

_BODY_FONT = "等线"
_HEADING_FONT = "黑体"


def _apply_run_font(run, *, name: str, size: float, bold: bool = False, gray: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if gray:
        run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    _apply_run_font(run, name=_HEADING_FONT, size=14, bold=True)


def _add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        _apply_run_font(run, name=_BODY_FONT, size=10.5)


def _add_numbered(document: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"{index}. {item}")
        _apply_run_font(run, name=_BODY_FONT, size=10.5)


def _add_action_table(document: Document, minutes: Minutes) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["序号", "待办事项", "责任人", "期限"]
    for cell, title in zip(table.rows[0].cells, headers, strict=False):
        cell.text = ""
        run = cell.paragraphs[0].add_run(title)
        _apply_run_font(run, name=_HEADING_FONT, size=10.5, bold=True)

    for index, item in enumerate(minutes.action_items, start=1):
        cells = table.add_row().cells
        values = [str(index), item.task, item.owner or "—", item.due or "—"]
        for cell, value in zip(cells, values, strict=False):
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            _apply_run_font(run, name=_BODY_FONT, size=10.5)


def build_docx(
    minutes: Minutes,
    *,
    include_transcript: bool = False,
    transcript: str | None = None,
) -> bytes:
    """把结构化纪要渲染成 Word 文档字节流。"""

    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = _BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), _BODY_FONT)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(minutes.title)
    _apply_run_font(run, name=_HEADING_FONT, size=18, bold=True)

    meta_parts: list[str] = []
    if minutes.meeting_date:
        meta_parts.append(f"会议日期：{minutes.meeting_date}")
    if minutes.participants:
        meta_parts.append(f"参会人：{'、'.join(minutes.participants)}")
    if meta_parts:
        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = meta.add_run("　".join(meta_parts))
        _apply_run_font(run, name=_BODY_FONT, size=10, gray=True)

    _add_heading(document, "一、会议摘要")
    paragraph = document.add_paragraph()
    run = paragraph.add_run(minutes.overview or "（无）")
    _apply_run_font(run, name=_BODY_FONT, size=10.5)

    _add_heading(document, "二、议题与讨论")
    if minutes.topics:
        for index, topic in enumerate(minutes.topics, start=1):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"{index}. {topic.title}")
            _apply_run_font(run, name=_BODY_FONT, size=10.5, bold=True)
            _add_bullets(document, topic.points)
    else:
        paragraph = document.add_paragraph("（无）")

    _add_heading(document, "三、决议事项")
    if minutes.decisions:
        _add_numbered(document, minutes.decisions)
    else:
        document.add_paragraph("（无）")

    _add_heading(document, "四、待办事项")
    if minutes.action_items:
        _add_action_table(document, minutes)
    else:
        document.add_paragraph("（无）")

    if minutes.risks:
        _add_heading(document, "五、风险与遗留问题")
        _add_bullets(document, minutes.risks)

    if minutes.next_steps:
        _add_heading(document, "六、下一步计划")
        _add_bullets(document, minutes.next_steps)

    if include_transcript and transcript:
        document.add_page_break()
        _add_heading(document, "附：会议转写全文")
        for line in transcript.splitlines():
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            _apply_run_font(run, name=_BODY_FONT, size=10.5)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
